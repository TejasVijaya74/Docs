import re
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    elif suffix == ".docx":
        return _extract_docx(file_path)
    elif suffix in (".png", ".jpg", ".jpeg"):
        from app.services.ocr import extract_text_from_image
        return extract_text_from_image(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        pages = []
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                pages.append(text)
        doc.close()
        raw = "\n\n".join(pages)
        if len(raw.strip()) > 100:
            return _clean_text(raw)
    except Exception as e:
        logger.warning(f"PyMuPDF failed: {e}, falling back to pdfplumber")

    try:
        import pdfplumber
        pages = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
        return _clean_text("\n\n".join(pages))
    except Exception as e:
        logger.error(f"pdfplumber also failed: {e}")
        raise RuntimeError("PDF text extraction failed with all available backends.")


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return _clean_text("\n\n".join(paragraphs))


def _clean_text(text: str) -> str:
    text = re.sub(r"\r\n|\r", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[^\S\n]+\n", "\n", text)
    return text.strip()